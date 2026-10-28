"""PDGM comorbidity adjustment (CY 2026), read from the CMS source file.

Rules, from the CY 2026 Low/High Comorbidity Adjustment file in
``data/cms/cy2026/comorbidity_interactions.xlsx`` (see SOURCES.md):

* **Low:** at least one secondary diagnosis is in one of the 20 low-comorbidity
  subgroups.
* **High:** two secondary diagnoses are in subgroups that form one of the 98
  CMS interaction pairs. High takes precedence over Low.
* Only secondary diagnoses count. The primary diagnosis's subgroup never
  contributes to the adjustment.

This replaces a hand-entered table that matched 2 of the 98 CMS interaction
pairs, modelled Low as a pair rather than a single subgroup, and counted the
primary diagnosis.

Subgroups are spelled ``Heart 11`` / ``Neoplasms 2`` in the CMS file and
``Heart_11`` / ``Neoplasm_2`` in the diagnosis mapping; everything here uses the
mapping spelling, and inputs in either spelling are accepted.
"""

import threading
from pathlib import Path
from typing import List

import openpyxl

DATA_FILE = (Path(__file__).resolve().parents[2]
             / 'data' / 'cms' / 'cy2026' / 'comorbidity_interactions.xlsx')

LOW_SUBGROUP_COUNT = 20
INTERACTION_COUNT = 98

_lock = threading.Lock()
_rules = None


def _norm(name):
    name = '_'.join(str(name).split())
    if name.startswith('Neoplasms_'):
        name = 'Neoplasm_' + name[len('Neoplasms_'):]
    return name


def load_rules():
    """Return ``{'low': {subgroup}, 'pairs': {frozenset: n}, 'descriptions': {subgroup: text}}``."""
    global _rules
    if _rules is not None:
        return _rules
    with _lock:
        if _rules is not None:
            return _rules
        # data_only: CMS fills the description columns with VLOOKUP formulas,
        # so read the cached values instead of the formula text.
        workbook = openpyxl.load_workbook(DATA_FILE, read_only=True, data_only=True)
        try:
            descriptions = {}
            low = set()
            for row in list(workbook['Low Comorbidity'].iter_rows(values_only=True))[2:]:
                if row[0]:
                    subgroup = _norm(row[0])
                    low.add(subgroup)
                    descriptions[subgroup] = (row[1] or '').strip()
            pairs = {}
            for row in list(workbook['High Comorbidity'].iter_rows(values_only=True))[2:]:
                if not row[0]:
                    continue
                a, b = _norm(row[1]), _norm(row[3])
                pairs[frozenset((a, b))] = int(row[0])
                descriptions.setdefault(a, (row[2] or '').strip())
                descriptions.setdefault(b, (row[4] or '').strip())
        finally:
            workbook.close()
        if len(low) != LOW_SUBGROUP_COUNT or len(pairs) != INTERACTION_COUNT:
            raise ValueError(
                'Expected %d low subgroups and %d interaction pairs in %s, parsed %d and %d'
                % (LOW_SUBGROUP_COUNT, INTERACTION_COUNT, DATA_FILE, len(low), len(pairs)))
        _rules = {'low': low, 'pairs': pairs, 'descriptions': descriptions}
        return _rules


def _subgroups(values):
    return sorted({_norm(v) for v in values if v and v != 'No_group'})


def explain_comorbidity_adjustment(secondary_subgroups: List[str]) -> dict:
    """The adjustment plus the low subgroups and interaction pairs behind it."""
    rules = load_rules()
    present = _subgroups(secondary_subgroups)
    low = [s for s in present if s in rules['low']]
    interactions = []
    for i, a in enumerate(present):
        for b in present[i + 1:]:
            number = rules['pairs'].get(frozenset((a, b)))
            if number:
                interactions.append({'interaction': number, 'subgroups': [a, b]})
    adjustment = 'High' if interactions else 'Low' if low else 'None'
    return {
        'adjustment': adjustment,
        'low_subgroups': low,
        'interactions': interactions,
    }


def determine_comorbidity_adjustment(
    primary_subgroup: str,
    secondary_subgroups: List[str],
) -> str:
    """Return 'None', 'Low' or 'High' for a period's secondary diagnoses.

    ``primary_subgroup`` is accepted for compatibility with existing callers
    and deliberately ignored: CMS bases the adjustment on secondary diagnoses
    only.
    """
    return explain_comorbidity_adjustment(secondary_subgroups)['adjustment']
